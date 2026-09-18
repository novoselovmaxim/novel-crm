import re
from pathlib import Path
from decimal import Decimal, InvalidOperation
from typing import Optional
from datetime import datetime as dt

import openpyxl
from sqlalchemy import select

from .database import async_session
from .models_ved import VedDeclaration, VedCompanyProfile


def parse_number(val) -> Optional[float]:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    s = re.sub(r"[^\d.\-]", "", s)
    if not s or s == "." or s == "-":
        return None
    try:
        return float(s)
    except (ValueError, InvalidOperation):
        return None


def safe_str(val) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    return s if s and s.lower() not in ("nan", "none", "", "#n/a") else None


def detect_format(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    cols = ws.max_column or 0
    first_cell = str(ws.cell(1, 1).value or "").lower()
    wb.close()

    if "направление" in first_cell and cols < 60:
        return "export_2025"
    if "№ п/п" in first_cell and cols > 100:
        return "export_aggregated"
    if "firm" in first_cell or "гтд" in first_cell.lower():
        if cols > 85:
            return "import_general"
        return "import_europe"
    if cols > 100:
        return "export_aggregated"
    return "import_europe"


def extract_key_fields(row_data: dict, fmt: str) -> dict:
    fields = {}

    if fmt in ("import_europe", "import_general"):
        inn_raw = (
            row_data.get("G081 (ИНН получателя)")
            or row_data.get("G091 (ИНН контрактодержателя)")
            or row_data.get("G141 (ИНН декларанта)")
        )
        fields["inn"] = str(int(inn_raw)) if isinstance(inn_raw, (int, float)) else safe_str(inn_raw)

        fields["company_name"] = safe_str(
            row_data.get("G082 (Наименование получателя)")
            or row_data.get("G092 (Наименование контрактодержателя)")
            or row_data.get("G142 (Наименование декларанта)")
        )

        fields["country_from"] = safe_str(
            row_data.get("G15A (Код страны отправления)")
            or row_data.get("G15 (Страна отправления)")
        )
        fields["country_to"] = safe_str(row_data.get("G17A (Код страны назначения)"))

        hs_raw = row_data.get("G33 (Код товара по ТН ВЭД)")
        fields["hs_code"] = str(int(hs_raw)) if isinstance(hs_raw, (int, float)) else safe_str(hs_raw)

        fields["customs_value"] = parse_number(row_data.get("G45 (Таможенная стоимость)"))
        fields["invoice_value"] = parse_number(row_data.get("G42 (Фактурная стоимость)"))
        fields["stat_value"] = parse_number(row_data.get("G46 (Статистическая стоимость, USD.)"))
        fields["net_weight"] = parse_number(row_data.get("G38 (Вес нетто, кг)"))
        fields["gross_weight"] = parse_number(row_data.get("G35 (Вес брутто, кг)"))
        fields["incoterms"] = safe_str(row_data.get("G202 (Условие поставки)"))

        date_str = safe_str(row_data.get("GD1 (Дата выпуска)"))
        if date_str:
            for fmt_d in ("%d.%m.%Y", "%Y-%m-%d"):
                try:
                    fields["declaration_date"] = dt.strptime(date_str, fmt_d).date()
                    break
                except ValueError:
                    continue

        firm_info = safe_str(row_data.get("FIRM (Доп.информация о контрактодержателе (Росстат))"))
        if firm_info:
            phone_m = re.search(r"Тел[^|]*?:\s*([^\|]+)", firm_info)
            email_m = re.search(r"[Ee]-?mail[^|]*?:\s*([^\|]+)", firm_info)
            site_m = re.search(r"Сайт[^|]*?:\s*([^\|]+)", firm_info)
            dir_m = re.search(r"Рук\.[^|]*?:\s*([^\|]+)", firm_info)
            if phone_m:
                fields["contact_phone"] = phone_m.group(1).strip()
            if email_m:
                fields["contact_email"] = email_m.group(1).strip()
            if site_m:
                fields["website"] = site_m.group(1).strip()
            if dir_m:
                fields["director"] = dir_m.group(1).strip()

    elif fmt == "export_2025":
        fields["inn"] = safe_str(row_data.get("ИНН отправителя"))
        fields["company_name"] = safe_str(row_data.get("Наименование отправителя"))
        fields["country_from"] = safe_str(row_data.get("Страна отправления"))
        fields["country_to"] = safe_str(row_data.get("Страна назначения"))

        hs_raw = row_data.get("Код ТН ВЭД. Знак звёздочки (*)...")
        fields["hs_code"] = str(int(hs_raw)) if isinstance(hs_raw, (int, float)) else safe_str(hs_raw)

        fields["customs_value"] = parse_number(row_data.get("Таможенная стоимость"))
        fields["invoice_value"] = parse_number(row_data.get("Фактурная стоимость"))
        fields["stat_value"] = parse_number(row_data.get("Статистическая стоимость"))
        fields["net_weight"] = parse_number(row_data.get("Вес нетто"))
        fields["gross_weight"] = parse_number(row_data.get("Вес брутто"))
        fields["incoterms"] = safe_str(row_data.get("Условие поставки в соответствии с Incoterms"))

        date_str = safe_str(row_data.get("Дата выпуска"))
        if date_str:
            for fmt_d in ("%Y-%m-%d", "%d.%m.%Y"):
                try:
                    fields["declaration_date"] = dt.strptime(date_str, fmt_d).date()
                    break
                except ValueError:
                    continue

        fields["contact_phone"] = safe_str(row_data.get("Телефоны"))
        fields["contact_email"] = safe_str(row_data.get("Почта"))
        fields["website"] = safe_str(row_data.get("Сайты"))
        fields["director"] = safe_str(row_data.get("Директор") or row_data.get("Гендиректор"))

    elif fmt == "export_aggregated":
        fields["inn"] = safe_str(row_data.get("ИНН декларанта"))
        if fields["inn"] and fields["inn"].replace(".", "").isdigit():
            fields["inn"] = str(int(float(fields["inn"])))
        fields["company_name"] = safe_str(
            row_data.get("Компания - наименование")
            or row_data.get("Наименование декларанта")
        )

        fields["country_to"] = safe_str(row_data.get("Страны назначения"))
        fields["customs_value"] = parse_number(row_data.get("Сумма статистической стоимости, USD"))
        fields["stat_value"] = parse_number(row_data.get("Сумма статистической стоимости, USD"))
        fields["net_weight"] = parse_number(row_data.get("Вес нетто, кг"))

        hs_raw = row_data.get("Коды ТН ВЭД")
        if hs_raw:
            codes = [c.strip() for c in str(hs_raw).split(",") if c.strip()]
            fields["hs_code"] = codes[0] if codes else None
        else:
            fields["hs_code"] = None

        fields["contact_phone"] = safe_str(row_data.get("Компания - телефон"))
        fields["contact_email"] = safe_str(row_data.get("Компания - e-mail"))
        fields["website"] = safe_str(row_data.get("Компания - сайт"))
        fields["director"] = safe_str(row_data.get("Компания - руководитель"))
        fields["address"] = safe_str(row_data.get("Компания - адрес"))
        fields["region"] = safe_str(row_data.get("Компания - регион регистрации"))
        fields["revenue"] = parse_number(row_data.get("Компания - выручка"))
        fields["employees"] = int(emp) if (emp := parse_number(row_data.get("Компания - Среднесписочная численность работников"))) else None
        fields["ogrn"] = safe_str(row_data.get("Компания - ОГРН"))
        fields["activity"] = safe_str(row_data.get("Компания - вид деятельности"))

        fields["direction"] = "export"

    return fields


def parse_ved_file(file_path: str) -> list[dict]:
    fmt = detect_format(file_path)
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    headers = []
    for col in range(1, (ws.max_column or 0) + 1):
        val = ws.cell(1, col).value
        headers.append(safe_str(val) or f"col_{col}")

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        row_data = {}
        for i, val in enumerate(row):
            if i < len(headers):
                row_data[headers[i]] = val
        rows.append(row_data)
    wb.close()

    result = []
    for row_data in rows:
        key_fields = extract_key_fields(row_data, fmt)
        if not key_fields.get("inn") and not key_fields.get("company_name"):
            continue
        key_fields["source_file"] = Path(file_path).name
        if "direction" not in key_fields:
            key_fields["direction"] = "import" if "import" in fmt else "export"
        key_fields["row_data"] = row_data
        result.append(key_fields)

    return result


async def import_ved_files(file_paths: list[str]) -> dict:
    all_records = []
    for fp in file_paths:
        try:
            records = parse_ved_file(fp)
            all_records.extend(records)
        except Exception as e:
            print(f"Error parsing {fp}: {e}")

    profiles_created = 0
    profiles_updated = 0
    declarations_inserted = 0

    async with async_session() as session:
        for rec in all_records:
            decl = VedDeclaration(
                source_file=rec["source_file"],
                direction=rec["direction"],
                inn=rec.get("inn"),
                company_name=rec.get("company_name"),
                country_from=rec.get("country_from"),
                country_to=rec.get("country_to"),
                hs_code=rec.get("hs_code"),
                customs_value=Decimal(str(rec["customs_value"])) if rec.get("customs_value") else None,
                invoice_value=Decimal(str(rec["invoice_value"])) if rec.get("invoice_value") else None,
                stat_value=Decimal(str(rec["stat_value"])) if rec.get("stat_value") else None,
                net_weight=Decimal(str(rec["net_weight"])) if rec.get("net_weight") else None,
                gross_weight=Decimal(str(rec["gross_weight"])) if rec.get("gross_weight") else None,
                incoterms=rec.get("incoterms"),
                declaration_date=rec.get("declaration_date"),
                row_data=rec["row_data"],
            )
            session.add(decl)
            declarations_inserted += 1

        await session.flush()

        inn_groups = {}
        for rec in all_records:
            inn = rec.get("inn")
            if not inn:
                continue
            if inn not in inn_groups:
                inn_groups[inn] = {
                    "inn": inn,
                    "company_name": rec.get("company_name"),
                    "directions": set(),
                    "declarations": 0,
                    "customs_value": Decimal("0"),
                    "stat_value": Decimal("0"),
                    "net_weight": Decimal("0"),
                    "dest_countries": set(),
                    "src_countries": set(),
                    "hs_codes": set(),
                    "source_files": set(),
                    "contact_phone": rec.get("contact_phone"),
                    "contact_email": rec.get("contact_email"),
                    "website": rec.get("website"),
                    "director": rec.get("director"),
                    "address": rec.get("address"),
                    "region": rec.get("region"),
                    "revenue": rec.get("revenue"),
                    "employees": rec.get("employees"),
                    "ogrn": rec.get("ogrn"),
                    "activity": rec.get("activity"),
                }
            g = inn_groups[inn]
            g["directions"].add(rec["direction"])
            g["declarations"] += 1
            if rec.get("customs_value"):
                g["customs_value"] += Decimal(str(rec["customs_value"]))
            if rec.get("stat_value"):
                g["stat_value"] += Decimal(str(rec["stat_value"]))
            if rec.get("net_weight"):
                g["net_weight"] += Decimal(str(rec["net_weight"]))
            if rec.get("country_to"):
                for c in rec["country_to"].split(","):
                    c = c.strip()
                    if c:
                        g["dest_countries"].add(c)
            if rec.get("country_from"):
                g["src_countries"].add(rec["country_from"])
            if rec.get("hs_code"):
                g["hs_codes"].add(rec["hs_code"])
            g["source_files"].add(rec["source_file"])

            for field in ("contact_phone", "contact_email", "website", "director", "address", "region", "revenue", "employees", "ogrn", "activity"):
                if rec.get(field) and not g.get(field):
                    g[field] = rec[field]

        for inn, g in inn_groups.items():
            result = await session.execute(
                select(VedCompanyProfile).where(VedCompanyProfile.inn == inn)
            )
            existing = result.scalar_one_or_none()

            dest_list = sorted(g["dest_countries"])
            src_list = sorted(g["src_countries"])
            hs_list = sorted(g["hs_codes"])[:20]
            sf_list = sorted(g["source_files"])
            dir_list = sorted(g["directions"])

            if existing:
                existing.company_name = g["company_name"] or existing.company_name
                existing.directions = dir_list
                existing.total_declarations = g["declarations"]
                existing.total_customs_value = g["customs_value"]
                existing.total_stat_value = g["stat_value"]
                existing.total_net_weight = g["net_weight"]
                existing.destination_countries = dest_list
                existing.source_countries = src_list
                existing.hs_codes = hs_list
                existing.source_files = sf_list
                if g.get("contact_phone"):
                    existing.contact_phone = g["contact_phone"]
                if g.get("contact_email"):
                    existing.contact_email = g["contact_email"]
                if g.get("website"):
                    existing.website = g["website"]
                if g.get("director"):
                    existing.director = g["director"]
                if g.get("address"):
                    existing.address = g["address"]
                if g.get("region"):
                    existing.region = g["region"]
                if g.get("revenue"):
                    existing.revenue = Decimal(str(g["revenue"]))
                if g.get("employees"):
                    existing.employees = g["employees"]
                if g.get("ogrn"):
                    existing.ogrn = g["ogrn"]
                if g.get("activity"):
                    existing.activity = g["activity"]
                profiles_updated += 1
            else:
                profile = VedCompanyProfile(
                    inn=inn,
                    company_name=g["company_name"],
                    directions=dir_list,
                    total_declarations=g["declarations"],
                    total_customs_value=g["customs_value"],
                    total_stat_value=g["stat_value"],
                    total_net_weight=g["net_weight"],
                    destination_countries=dest_list,
                    source_countries=src_list,
                    hs_codes=hs_list,
                    source_files=sf_list,
                    contact_phone=g.get("contact_phone"),
                    contact_email=g.get("contact_email"),
                    website=g.get("website"),
                    director=g.get("director"),
                    address=g.get("address"),
                    region=g.get("region"),
                    revenue=Decimal(str(g["revenue"])) if g.get("revenue") else None,
                    employees=g.get("employees"),
                    ogrn=g.get("ogrn"),
                    activity=g.get("activity"),
                )
                session.add(profile)
                profiles_created += 1

        await session.commit()

    return {
        "profiles_created": profiles_created,
        "profiles_updated": profiles_updated,
        "declarations_inserted": declarations_inserted,
    }
